from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.chart.data import ChartData
from pptx.enum.dml import MSO_THEME_COLOR
from pathlib import Path

OUT = Path(__file__).parent
PPTX = OUT / "FlowGuard_rebuilt_deck.pptx"

W, H = 13.333, 7.5
CREAM = RGBColor(247, 240, 230)
INK = RGBColor(42, 37, 35)
MUTED = RGBColor(99, 88, 82)
TERRA = RGBColor(179, 78, 58)
SAGE = RGBColor(65, 108, 91)
PALE = RGBColor(239, 228, 213)
WHITE = RGBColor(255,255,255)

prs = Presentation()
prs.slide_width = Inches(W); prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]

def box(slide, x, y, w, h, fill=None, line=None, radius=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill or CREAM
    sh.line.color.rgb = line or fill or CREAM
    return sh

def text(slide, txt, x, y, w, h, size=18, color=INK, bold=False, font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.clear(); tf.word_wrap = True; tf.margin_left = Inches(.04); tf.margin_right = Inches(.04)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = txt; r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb

def rich(slide, runs, x, y, w, h, size=18, color=INK, font="Aptos"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=tb.text_frame; tf.clear(); tf.word_wrap=True
    p=tf.paragraphs[0]
    for s,b,c in runs:
        r=p.add_run(); r.text=s; r.font.name=font; r.font.size=Pt(size); r.font.bold=b; r.font.color.rgb=c or color
    return tb

def base(slide, num, kicker, title, subtitle=None):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = CREAM
    box(slide, .45, .38, .12, .72, fill=TERRA)
    text(slide, kicker.upper(), .72, .42, 4.8, .25, 10, TERRA, True, "Aptos")
    text(slide, title, .72, .72, 11.9, .72, 28, INK, True, "Georgia")
    if subtitle: text(slide, subtitle, .74, 1.48, 11.7, .38, 12, MUTED, False)
    text(slide, f"FLOWGUARD  /  {num:02d}", 11.15, 7.08, 1.7, .2, 8, MUTED, True, "Aptos", PP_ALIGN.RIGHT)

def foot(slide, txt): text(slide, txt, .74, 6.78, 11.1, .22, 8, MUTED)
def note(slide, txt):
    try:
        ns = slide.notes_slide
        ns.notes_text_frame.text = txt
    except Exception: pass

def pill(slide, txt, x, y, w, fill=PALE, color=INK):
    box(slide,x,y,w,.32,fill=fill,line=fill,radius=True); text(slide,txt,x+.08,y+.055,w-.16,.2,9,color,True)

slides = []

# 1
s=prs.slides.add_slide(blank); slides.append(s); base(s,1,"Purpose","Make feasibility a first-class search signal","A controlled physical-design experiment about where EDA time goes.")
text(s,"FlowGuard asks a practical question: when every RTL-to-GDS run costs real minutes, can a search policy spend fewer of them on configurations that were never physically viable?",.9,2.3,7.1,1.35,20,INK,False,"Georgia")
box(s,8.65,2.18,3.65,2.35,fill=INK,line=INK,radius=True)
text(s,"EI × P(feasible)",8.98,2.58,3.0,.5,25,CREAM,True,"Georgia",PP_ALIGN.CENTER)
text(s,"same budget  /  different allocation",9.02,3.35,2.9,.35,12,PALE,False,"Aptos",PP_ALIGN.CENTER)
text(s,"The claim is about failed-run compute, not a promise of better final QoR.",.92,4.85,7.1,.5,14,TERRA,True,"Aptos")
foot(s,"Scope: crossbar-v2 / fixed 19.9 ns / 36-config oracle")
note(s,"Set the frame: this is a resource-allocation question. The design, flow, and evaluation budget stay fixed; only the acquisition policy changes.")

# 2
s=prs.slides.add_slide(blank); slides.append(s); base(s,2,"Boundary","One crossbar. Two physical outcomes.","The 19.9 ns crossbar-v2 boundary is a 36-configuration oracle, not the older stress benchmark.")
text(s,"Same crossbar, fixed constraints, different outcomes: six feasible configurations and thirty failures.",.9,2.05,11.1,.55,18,INK,False,"Georgia")
for i,(title,sub,fill,accent) in enumerate([("FEASIBLE REGION","6 / 36 · positive checks",SAGE,CREAM),("FAILURE REGION","30 / 36 · recorded failures",RGBColor(247,225,216),TERRA)]):
    x=.95+i*6.1
    box(s,x,3.05,5.35,1.95,fill=fill,line=accent,radius=True)
    text(s,title,x+.35,3.4,4.65,.35,14,accent,True)
    text(s,sub,x+.35,3.95,4.65,.3,14,CREAM if i == 0 else INK,False)
    text(s,"candidate settings → physical result",x+.35,4.42,4.65,.25,10,PALE if i == 0 else MUTED)
box(s,6.25,3.65,.82,.12,fill=TERRA,line=TERRA); box(s,6.98,3.47,.12,.48,fill=TERRA,line=TERRA)
text(s,"19.9 ns",5.65,5.45,2.1,.35,24,TERRA,True,"Georgia",PP_ALIGN.CENTER)
foot(s,"Authoritative crossbar-v2 oracle: 36 configs, 6 feasible, 30 failed. Older stress evidence is excluded.")
note(s,"The crossbar oracle is the evidence here: 36 configurations at fixed 19.9 ns, with 6 feasible and 30 failed. Do not import older stress characterization.")

# 3
s=prs.slides.add_slide(blank); slides.append(s); base(s,3,"Model","Quality plus feasibility, separated on purpose.","FlowGuard learns what is good only where a run is physically valid—and learns validity from every attempted trial.")
box(s,.9,2.25,4.6,2.75,fill=WHITE,line=PALE,radius=True); text(s,"QUALITY MODEL",1.25,2.62,3.9,.3,12,TERRA,True); text(s,"GP → expected improvement",1.25,3.12,3.9,.38,20,INK,True,"Georgia"); text(s,"trained on feasible trials only",1.25,3.78,3.9,.35,13,MUTED)
box(s,7.85,2.25,4.6,2.75,fill=WHITE,line=PALE,radius=True); text(s,"FEASIBILITY MODEL",8.2,2.62,3.9,.3,12,SAGE,True); text(s,"RF → P(feasible)",8.2,3.12,3.9,.38,20,INK,True,"Georgia"); text(s,"trained on all attempted trials",8.2,3.78,3.9,.35,13,MUTED)
box(s,5.58,3.25,2.15,.18,fill=TERRA,line=TERRA); text(s,"×",6.18,3.57,.95,.65,32,TERRA,True,"Georgia",PP_ALIGN.CENTER)
text(s,"acquisition = EI × P(feasible)",3.45,5.55,6.5,.48,22,INK,True,"Georgia",PP_ALIGN.CENTER)
foot(s,"Production objective qor_v1: 0.5 critical delay + 0.3 wirelength + 0.2 cell area (normalized).")
note(s,"The product is deliberately simple: expected improvement is discounted when the feasibility model says the physical outcome is unlikely. This is not a hard rejection rule.")

# 4
s=prs.slides.add_slide(blank); slides.append(s); base(s,4,"Evidence","One decision. One authentic physical artifact.","Replay selects cb36_002; the layout image comes from its separate banked physical run.")
box(s,.75,2.12,4.55,3.35,fill=WHITE,line=PALE,radius=True); s.shapes.add_picture(str(OUT/"assets"/"cb36_002.png"), Inches(1.0), Inches(2.48), width=Inches(4.05), height=Inches(1.25)); text(s,"cb36_002 · SHA-256 35fbbd7e…a12dc5de",1.0,3.95,4.05,.3,9,MUTED); text(s,"separate-run artifact",1.0,4.55,4.05,.3,13,TERRA,True)
box(s,5.65,2.12,6.8,3.35,fill=INK,line=INK,radius=True); text(s,"FlowGuard-Calibrated · seed 11 · call 10",6.02,2.48,6.0,.3,13,CREAM,True); text(s,"selected cb36_002",6.02,2.95,5.8,.4,24,CREAM,True,"Georgia"); text(s,"EI 0.0 × P(feasible) 0.08 · rank 1 / 27\nfeasible · QoR −0.093068 · setup +0.190903 ns\nDRC 0 · routing 100 · LVS/signoff pass",6.02,3.62,5.9,1.1,13,PALE)
foot(s,"Replay launched zero physical runs; PNG is from earlier trial_cb36_002. SHA-256 verified against 06_demo.json.")
note(s,"Show the authentic PNG and keep provenance separate: the replay is offline and launches zero EDA processes. The image is attached by candidate ID from an earlier banked physical run, with the exact SHA recorded in 06_demo.json.")

# 5
s=prs.slides.add_slide(blank); slides.append(s); base(s,5,"Protocol","Freeze the comparison before interpreting it.","A matched 24-evaluation replay across 13 seeds and three methods.")
rows=[("Design","crossbar_datapath_v2"),("Clock","19.9 ns (fixed)"),("Pool","36 configurations · 6 / 30 oracle"),("Budget","24 evaluations: 8 shared init + 16 adaptive"),("Methods","Vanilla · FlowGuard · EI-only")]
table=s.shapes.add_table(len(rows),2,Inches(1.0),Inches(2.25),Inches(11.2),Inches(3.2)).table; table.columns[0].width= Inches(3.2); table.columns[1].width= Inches(8.0)
for r,(a,b) in enumerate(rows):
    for c,v in enumerate((a,b)):
        cell=table.cell(r,c); cell.text=v; cell.fill.solid(); cell.fill.fore_color.rgb=INK if c==0 else WHITE
        for p in cell.text_frame.paragraphs:
            p.runs[0].font.name="Aptos"; p.runs[0].font.size=Pt(14); p.runs[0].font.color.rgb=CREAM if c==0 else INK; p.runs[0].font.bold=(c==0)
foot(s,"No replay EDA: the bundle replays committed outcomes; physical metrics are banked oracle results.")
note(s,"The crossbar protocol is fixed at 19.9 ns with 36 configurations. Each method gets 24 evaluations, including the same eight shared initializations, across 13 seeds. Replay launches no EDA.")

# 6
s=prs.slides.add_slide(blank); slides.append(s); base(s,6,"Headline result","Feasibility weighting reduced failed-run compute by 4.1%","FlowGuard versus matched EI-only search across 13 replay seeds")
text(s,"10.28 fewer minutes spent on failed configurations per replay, on average",.95,1.95,11.5,.45,20,TERRA,True,"Georgia",PP_ALIGN.CENTER)
cd=ChartData(); cd.categories=["EI-only","FlowGuard-Calibrated"]; cd.add_series("Mean failed EDA cost (min)",[249.91,239.63])
chart=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(1.0),Inches(2.55),Inches(6.8),Inches(3.25),cd).chart
chart.has_legend=False; chart.value_axis.minimum_scale=0; chart.value_axis.maximum_scale=270; chart.value_axis.major_unit=50; chart.value_axis.has_major_gridlines=True
chart.category_axis.reverse_order=True; chart.value_axis.tick_labels.font.size=Pt(10); chart.category_axis.tick_labels.font.size=Pt(11)
chart.series[0].format.fill.solid(); chart.series[0].format.fill.fore_color.rgb=TERRA
box(s,8.45,2.68,3.85,2.45,fill=INK,line=INK,radius=True); text(s,"−4.11%",8.82,3.05,3.1,.62,34,CREAM,True,"Georgia",PP_ALIGN.CENTER); text(s,"derived: (14994.8−14378.0) / 14994.8",8.75,3.86,3.2,.28,10,PALE,False,"Aptos",PP_ALIGN.CENTER); text(s,"same final QoR and best-call outcome; total cost slightly higher",8.82,4.35,3.1,.55,11,CREAM,False,"Aptos",PP_ALIGN.CENTER)
foot(s,"Derived from verified failed means 14994.8 s and 14378.0 s: 616.8 s = 10.28 min. Editable zero-based chart; sequential EDA-runtime estimate.")
note(s,"Click the chart and show the editable data. The 10.28 minutes is (14994.8−14378.0) seconds divided by 60. The 4.11 percent is 616.8 divided by 14994.8. All methods tie final QoR; FlowGuard matches EI-only on best-call across all 13 seeds.")

# 7
s=prs.slides.add_slide(blank); slides.append(s); base(s,7,"Support comparison","Three methods. Vanilla spends least.","Across 13 seeds, all three tie final QoR; cost separates failed and total runtime means.")
headers=["Method","Mean feasible","Failed EDA mean (s)","Total EDA mean (s)","Final QoR","Best-call"]
data=[("Vanilla","6.00","13046.8","17230.7","−0.093068","baseline"),("FlowGuard","4.38","14378.0","17307.3","−0.093068","= EI-only"),("EI-only","3.38","14994.8","17272.8","−0.093068","= FlowGuard")]
table=s.shapes.add_table(4,6,Inches(.85),Inches(2.15),Inches(11.7),Inches(2.45)).table
widths=[2.35,1.65,1.55,2.0,2.0,2.15]
for i,w in enumerate(widths): table.columns[i].width=Inches(w)
for c,h in enumerate(headers):
    cell=table.cell(0,c); cell.text=h; cell.fill.solid(); cell.fill.fore_color.rgb=INK
    cell.text_frame.paragraphs[0].runs[0].font.color.rgb=CREAM; cell.text_frame.paragraphs[0].runs[0].font.bold=True; cell.text_frame.paragraphs[0].runs[0].font.size=Pt(10)
for r,row in enumerate(data,1):
    for c,v in enumerate(row):
        cell=table.cell(r,c); cell.text=v; cell.fill.solid(); cell.fill.fore_color.rgb=WHITE if r%2 else PALE
        rr=cell.text_frame.paragraphs[0].runs[0]; rr.font.name="Aptos"; rr.font.size=Pt(12); rr.font.color.rgb=INK; rr.font.bold=(c==0)
box(s,1.0,5.05,11.3,.9,fill=RGBColor(235,245,235),line=SAGE,radius=True); text(s,"Vanilla is lowest on both failed EDA mean and total EDA mean; final QoR ties three ways.",1.3,5.3,10.7,.35,15,SAGE,True,"Aptos",PP_ALIGN.CENTER)
foot(s,"Means from final_evidence_bundle/03_recomputed_aggregates.json; feasible counts are seed means.")
note(s,"This is the complete crossbar support table. Vanilla has the lowest failed and total EDA means, while the primary endpoint is tied. FlowGuard and EI-only reach the best call identically across all 13 seeds.")

# 8
s=prs.slides.add_slide(blank); slides.append(s); base(s,8,"Reproducibility","Reproduce the decision, not just the layout.","Crossbar replay is an oracle/cache process, not a new physical campaign.")
box(s,.9,2.25,5.45,2.7,fill=WHITE,line=PALE,radius=True); text(s,"Older TPE / layout image",1.28,2.62,4.6,.3,13,TERRA,True); text(s,"physical artifact ≠ decision record",1.28,3.22,4.6,.4,18,INK,True,"Georgia"); text(s,"keep separate from crossbar evidence",1.28,4.05,4.6,.35,12,MUTED)
box(s,6.98,2.25,5.45,2.7,fill=INK,line=INK,radius=True); text(s,"FlowGuard / decision record",7.36,2.62,4.6,.3,13,CREAM,True); text(s,"seed · call · hash · scores · rank",7.36,3.22,4.6,.4,18,CREAM,True,"Georgia"); text(s,"cb36_002 replay record is traceable",7.36,4.05,4.6,.35,12,PALE)
text(s,"Limitations: no standalone combined-13 verification log; replay/oracle outcomes attach banked physical runs; older stress evidence is excluded.",1.0,5.55,11.3,.55,13,MUTED,False,"Aptos",PP_ALIGN.CENTER)
foot(s,"Bundle facts: 07_FACTS.md and VALIDATION.md. Physical image provenance is explicit on slide 4.")
note(s,"Separate the older TPE layout image from the FlowGuard decision record. The authoritative crossbar result is a replay/oracle process with banked physical outcomes. There is no standalone combined-13 verification log, and older stress evidence is excluded.")

# 9
s=prs.slides.add_slide(blank); slides.append(s); base(s,9,"Invitation","Reproduce the replay. Then run the hardware.","The crossbar bundle makes the decision traceable; the next step is independent physical confirmation.")
text(s,"Run `python3 experiments/replay_demo.py`, inspect cb36_002, then verify the banked artifact independently.",1.0,2.25,7.2,1.15,21,INK,False,"Georgia")
box(s,8.75,2.2,3.45,2.35,fill=TERRA,line=TERRA,radius=True); text(s,"REPRODUCE",9.1,2.7,2.75,.35,15,CREAM,True,"Aptos",PP_ALIGN.CENTER); text(s,"then falsify",9.1,3.35,2.75,.5,24,CREAM,True,"Georgia",PP_ALIGN.CENTER)
for i,t in enumerate(["1  freeze","2  run","3  inspect"]): pill(s,t,1.08+i*2.0,4.65,1.62,PALE,INK)
text(s,"Next test: independently reproduce the 13-seed tie and validate the separate physical artifact chain.",1.0,5.45,11.2,.5,15,TERRA,True,"Aptos",PP_ALIGN.CENTER)
foot(s,"Command from the crossbar-v2 replay workflow; no replay EDA is launched.")
note(s,"Close with the exact demo command: python3 experiments/replay_demo.py. It reproduces the trace, not a new physical run. Invite independent hardware confirmation of the separate artifact.")

prs.save(PPTX)

print(PPTX)
